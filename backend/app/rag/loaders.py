"""Document text extraction for the supported upload formats.

Each loader returns an :class:`Extraction` describing what happened, so the API
can report a real per-document status instead of pretending every upload was
indexed.  A PDF that cannot yield text is reported as ``failed`` with a reason,
never silently indexed as empty.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from ..core.logging import get_logger

logger = get_logger("app.rag.loaders")

TEXT_SUFFIXES = {".md", ".markdown", ".txt"}
PDF_SUFFIX = ".pdf"
DOCX_SUFFIX = ".docx"

# Encodings tried in order.  GB18030 covers legacy Chinese Windows exports.
_TEXT_ENCODINGS = ("utf-8", "utf-8-sig", "gb18030", "latin-1")

TYPE_LABELS = {
    ".md": "Markdown",
    ".markdown": "Markdown",
    ".txt": "Text",
    ".pdf": "PDF",
    ".docx": "Word",
}


@dataclass
class Extraction:
    """Result of reading one file."""

    text: str = ""
    ok: bool = False
    note: str = ""
    pages: int = 0
    error: str = ""
    warnings: list[str] = field(default_factory=list)


def type_label(suffix: str) -> str:
    """Human label for a file suffix."""
    return TYPE_LABELS.get(suffix.lower(), (suffix or "").lstrip(".").upper() or "Unknown")


# ---------------------------------------------------------------------------
# Plain text
# ---------------------------------------------------------------------------


def _read_text_file(path: Path) -> Extraction:
    for encoding in _TEXT_ENCODINGS:
        try:
            content = path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
        except OSError as exc:
            return Extraction(ok=False, error=f"读取失败：{exc}", note="文件无法读取")
        if content.strip():
            note = "" if encoding in {"utf-8", "utf-8-sig"} else f"以 {encoding} 编码解码"
            return Extraction(text=content, ok=True, note=note)
        return Extraction(ok=True, text="", note="文件内容为空")

    return Extraction(ok=False, error="文件编码无法识别", note="请转换为 UTF-8 后重试")


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def _read_pdf_file(path: Path) -> Extraction:
    try:
        from pypdf import PdfReader
    except ImportError:
        return Extraction(
            ok=False,
            error="缺少 pypdf 依赖",
            note="安装 pypdf 后即可解析 PDF 全文",
        )

    try:
        reader = PdfReader(str(path))
        if getattr(reader, "is_encrypted", False):
            try:
                reader.decrypt("")
            except Exception:
                return Extraction(ok=False, error="PDF 已加密", note="加密 PDF 暂不支持解析")

        pages = len(reader.pages)
        chunks: list[str] = []
        empty_pages = 0
        for page_number, page in enumerate(reader.pages, start=1):
            try:
                page_text = page.extract_text() or ""
            except Exception as exc:  # pragma: no cover - malformed page
                logger.warning("PDF page %d failed to extract: %s", page_number, exc)
                page_text = ""
            if page_text.strip():
                chunks.append(page_text)
            else:
                empty_pages += 1
    except Exception as exc:
        return Extraction(ok=False, error=f"PDF 解析失败：{exc}", note="文件可能已损坏")

    text = "\n\n".join(chunks).strip()
    warnings: list[str] = []
    if empty_pages:
        warnings.append(f"{empty_pages}/{pages} 页未提取到文本层（可能是扫描件）")

    if not text:
        return Extraction(
            ok=False,
            pages=pages,
            error="PDF 未提取到文本",
            note="该 PDF 很可能是扫描图片，需要 OCR，当前版本不支持",
            warnings=warnings,
        )

    note = f"已解析 {pages} 页文本"
    if warnings:
        note += f"；{warnings[0]}"
    return Extraction(text=text, ok=True, note=note, pages=pages, warnings=warnings)


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def _read_docx_file(path: Path) -> Extraction:
    try:
        import docx  # python-docx
    except ImportError:
        return Extraction(
            ok=False,
            error="缺少 python-docx 依赖",
            note="安装 python-docx 后即可解析 Word 全文",
        )

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        return Extraction(ok=False, error=f"DOCX 解析失败：{exc}", note="文件可能已损坏")

    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = (paragraph.text or "").strip()
        if not text:
            continue
        style = (getattr(paragraph, "style", None) and paragraph.style.name) or ""
        if style.lower().startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "2"
            lines.append(f"{'#' * min(int(level), 6)} {text}")
        else:
            lines.append(text)

    table_count = 0
    for table in document.tables:
        table_count += 1
        rows: list[str] = []
        for row in table.rows:
            cells = [(cell.text or "").strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                rows.append("| " + " | ".join(cells) + " |")
        if rows:
            header = rows[0]
            separator = "| " + " | ".join(["---"] * (header.count("|") - 1)) + " |"
            lines.append("")
            lines.extend([header, separator, *rows[1:]])
            lines.append("")

    text = "\n\n".join(line for line in lines if line is not None).strip()
    if not text:
        return Extraction(ok=False, error="DOCX 未提取到文本", note="文档正文为空")

    note = "已解析 Word 正文"
    if table_count:
        note += f"，含 {table_count} 个表格"
    return Extraction(text=text, ok=True, note=note)


# ---------------------------------------------------------------------------
# Dispatch
# ---------------------------------------------------------------------------


def extract_text(path: str | Path) -> Extraction:
    """Extract plain text from ``path`` based on its suffix."""
    path = Path(path)
    if not path.exists():
        return Extraction(ok=False, error="文件不存在", note="")
    if not path.is_file():
        return Extraction(ok=False, error="目标不是文件", note="")

    suffix = path.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        return _read_text_file(path)
    if suffix == PDF_SUFFIX:
        return _read_pdf_file(path)
    if suffix == DOCX_SUFFIX:
        return _read_docx_file(path)

    return Extraction(
        ok=False,
        error=f"不支持的文件类型 {suffix or '(无扩展名)'}",
        note="当前支持 .md / .txt / .pdf / .docx",
    )
