"""Presales solution generation and export.

Flow (identical whether triggered from Solution Studio or the Agent):

    requirement text → analyze_requirement → retrieve evidence
                     → grounded prompt → parse the eight mandated sections
                     → citations + markdown (+ optional DOCX export)

The eight sections are **enforced**, not suggested: whatever the model omits is
reported as missing rather than silently dropped, so a proposal never looks
complete when it is not.
"""

from __future__ import annotations

import re

from ..core.config import Settings, get_settings
from ..core.logging import get_logger
from ..core.tracing import TraceRecorder
from ..rag.citations import (
    analyse_grounding,
    build_citations,
    build_sources,
    format_context,
    strip_invalid_citations,
)
from ..rag.index import KnowledgeIndex, get_index
from ..rag.llm import get_llm_client
from ..rag.pipeline import RagPipeline
from ..rag.prompts import get_prompt_library
from ..schemas.rag import RetrievedChunk, UsageInfo
from ..schemas.solution import (
    RequirementAnalysis,
    SolutionResponse,
    SolutionSection,
)
from .requirement_service import analyze_requirement, build_requirement_text

logger = get_logger("app.services.solution")

SECTION_TITLES: tuple[tuple[str, str], ...] = (
    ("executive_summary", "Executive Summary"),
    ("customer_needs", "Customer Needs"),
    ("recommended_architecture", "Recommended Architecture"),
    ("product_mapping", "Product Mapping"),
    ("implementation_plan", "Implementation Plan"),
    ("case_references", "Case References"),
    ("risks", "Risks"),
    ("next_steps", "Next Steps"),
)

_HEADING = re.compile(r"^##\s+(.*?)\s*$", re.MULTILINE)
_HEADING_ALIASES: dict[str, str] = {
    "executive summary": "executive_summary",
    "客户背景理解": "executive_summary",
    "摘要": "executive_summary",
    "customer needs": "customer_needs",
    "核心需求与痛点": "customer_needs",
    "客户需求": "customer_needs",
    "recommended architecture": "recommended_architecture",
    "推荐解决方案": "recommended_architecture",
    "推荐架构": "recommended_architecture",
    "product mapping": "product_mapping",
    "功能模块设计": "product_mapping",
    "产品映射": "product_mapping",
    "implementation plan": "implementation_plan",
    "实施步骤建议": "implementation_plan",
    "实施计划": "implementation_plan",
    "case references": "case_references",
    "案例参考": "case_references",
    "risks": "risks",
    "风险": "risks",
    "next steps": "next_steps",
    "下一步": "next_steps",
}


def parse_sections(markdown: str) -> dict[str, str]:
    """Split the model output into the mandated sections."""
    text = markdown or ""
    matches = list(_HEADING.finditer(text))
    sections: dict[str, str] = {}

    for position, match in enumerate(matches):
        title = match.group(1).strip().lower()
        key = _HEADING_ALIASES.get(title)
        if key is None:
            for alias, alias_key in _HEADING_ALIASES.items():
                if alias in title:
                    key = alias_key
                    break
        if key is None:
            continue
        end = matches[position + 1].start() if position + 1 < len(matches) else len(text)
        body = text[match.end() : end].strip()
        if body:
            sections[key] = body

    return sections


def build_sections(parsed: dict[str, str]) -> tuple[list[SolutionSection], list[str]]:
    """Order the parsed sections and report which ones are missing."""
    sections: list[SolutionSection] = []
    missing: list[str] = []

    for key, title in SECTION_TITLES:
        body = parsed.get(key, "").strip()
        if not body:
            missing.append(title)
            body = "> 本次生成未产出该章节内容，建议人工补充或调整检索资料后重新生成。"
        sections.append(
            SolutionSection(
                key=key,  # type: ignore[arg-type]
                title=title,
                content=body,
                grounded=bool(re.search(r"\[\d+\]", body)),
            )
        )
    return sections, missing


def render_markdown(
    requirement: str,
    analysis: RequirementAnalysis,
    sections: list[SolutionSection],
    sources_markdown: str,
) -> str:
    """Render the downloadable proposal."""
    lines = [
        "# 售前解决方案（初稿）",
        "",
        "> 本方案由 Enterprise RAG Copilot 基于企业知识库资料生成，"
        "章节内容均标注了引用来源编号，未检索到依据的部分已显式说明。",
        "",
        "## 客户需求原文",
        "",
        requirement.strip() or "（未提供）",
        "",
        "## 需求解析",
        "",
        f"- 客户类型：{analysis.customer_type or '未提及'}",
        f"- 所属行业：{analysis.industry or '未提及'}",
        f"- 业务场景：{analysis.scenario or '未提及'}",
        f"- 解析来源：{'模型解析' if analysis.generated_by == 'llm' else '规则提取'}",
        "",
        "### 核心需求",
        *([f"- {item}" for item in analysis.core_needs] or ["- 待补充"]),
        "",
        "### 主要痛点",
        *([f"- {item}" for item in analysis.pain_points] or ["- 待补充"]),
        "",
        "### 关键约束",
        *([f"- {item}" for item in analysis.constraints] or ["- 暂无明确约束"]),
        "",
        "---",
        "",
    ]
    for section in sections:
        lines.extend([f"## {section.title}", "", section.content, ""])
    if sources_markdown:
        lines.extend(["---", "", "## 参考资料来源", "", sources_markdown, ""])
    return "\n".join(lines)


def _sources_markdown(sources) -> str:
    if not sources:
        return ""
    return "\n".join(
        f"- [{index}] {source.document_name}（{source.chunk_count} 处引用）"
        for index, source in enumerate(sources, start=1)
    )


def _failed_response(requirement: str, analysis: RequirementAnalysis, message: str) -> SolutionResponse:
    return SolutionResponse(
        requirement_text=requirement,
        analysis=analysis,
        sections=[],
        markdown=f"# 售前解决方案（生成失败）\n\n{message}\n",
        grounded=False,
        warnings=[message],
    )


def compose_solution(
    *,
    requirement: str,
    analysis: RequirementAnalysis,
    chunks: list[RetrievedChunk],
    settings: Settings | None = None,
    trace: TraceRecorder | None = None,
) -> SolutionResponse:
    """Generate the proposal from already-retrieved evidence."""
    settings = settings or get_settings()
    trace = trace or TraceRecorder()

    if not chunks:
        message = (
            "知识库中没有检索到可用于方案生成的资料。"
            "请先在「文档管理」中补充产品、行业方案与案例类文档，再重新生成。"
        )
        trace.add("generate", "生成方案", status="warning", summary=message)
        return _failed_response(requirement, analysis, message)

    context = format_context(chunks)
    citations = build_citations(chunks)

    with trace.step(
        "generate",
        "生成售前方案",
        tool="llm",
        inputs={"model": settings.llm_model, "context_blocks": len(chunks)},
    ) as step:
        client = get_llm_client(settings)
        if not client.configured:
            message = "未配置模型 API Key，无法生成方案正文。检索到的资料已在下方列出。"
            step.status = "failed"
            step.summary = message
            response = _failed_response(requirement, analysis, message)
            response.retrieved_chunks = chunks
            response.citations = citations
            response.sources = build_sources(citations)
            return response

        try:
            system, user = get_prompt_library(settings).render(
                "solution_generation",
                requirement=requirement,
                analysis=analysis.model_dump_json(indent=2),
                context=context,
                context_count=len(chunks),
            )
        except Exception as exc:
            message = f"Prompt 模板加载失败：{exc}"
            step.status = "failed"
            step.summary = message
            return _failed_response(requirement, analysis, message)

        completion = client.complete(system=system, user=user)
        if not completion.ok:
            step.status = "failed"
            step.summary = completion.error
            response = _failed_response(requirement, analysis, completion.error)
            response.retrieved_chunks = chunks
            response.citations = citations
            response.sources = build_sources(citations)
            return response

        step.summary = f"生成 {len(completion.text)} 字方案，消耗 {completion.total_tokens} tokens"
        step.outputs = {"chars": len(completion.text), "tokens": completion.total_tokens}

    # -- grounding: drop citation markers that point outside the supplied context
    groundings = analyse_grounding(completion.text, len(citations))
    cleaned = strip_invalid_citations(completion.text, len(citations))
    parsed = parse_sections(cleaned)
    sections, missing = build_sections(parsed)

    warnings: list[str] = []
    if groundings.invalid_indexes:
        warnings.append(
            f"方案中出现了 {len(groundings.invalid_indexes)} 个越界引用编号，已自动移除。"
        )
    if missing:
        warnings.append("以下章节未生成内容：" + "、".join(missing))
    if not groundings.has_citation:
        warnings.append("方案正文没有引用编号，溯源强度较弱，建议人工核对。")
    if analysis.generated_by == "rules":
        warnings.append("需求解析由规则引擎完成（未使用模型），建议人工复核。")

    final_citations = (
        [item for item in citations if item.index in set(groundings.cited_indexes)]
        if groundings.cited_indexes
        else citations
    )
    final_sources = build_sources(final_citations)
    markdown = render_markdown(requirement, analysis, sections, _sources_markdown(final_sources))

    return SolutionResponse(
        requirement_text=requirement,
        analysis=analysis,
        sections=sections,
        citations=final_citations,
        sources=final_sources,
        retrieved_chunks=chunks,
        trace=trace.steps,
        markdown=markdown,
        grounded=groundings.has_citation and not groundings.invalid_indexes,
        warnings=warnings,
        usage=UsageInfo(
            prompt_tokens=completion.prompt_tokens,
            completion_tokens=completion.completion_tokens,
            total_tokens=completion.total_tokens,
        ),
        model=completion.model,
    )


def generate_solution(
    requirement: str = "",
    *,
    form=None,
    settings: Settings | None = None,
    index: KnowledgeIndex | None = None,
    top_k: int | None = None,
    trace: TraceRecorder | None = None,
) -> SolutionResponse:
    """Full Solution Studio flow: analyse → retrieve → generate."""
    settings = settings or get_settings()
    index = index or get_index(settings)
    trace = trace or TraceRecorder()
    requirement_text = build_requirement_text(requirement, form)

    with trace.step("understand", "解析客户需求", tool="analyze_requirements") as step:
        analysis = analyze_requirement(requirement_text, settings=settings)
        step.summary = (
            f"客户类型 {analysis.customer_type or '未提及'}｜行业 {analysis.industry or '未提及'}"
            f"｜核心需求 {len(analysis.core_needs)} 条｜待确认 {len(analysis.missing_info)} 项"
        )
        step.outputs = analysis.model_dump()

    pipeline = RagPipeline(index, settings)
    query = analysis.search_query or requirement_text or "行业解决方案"

    with trace.step("retrieve", "检索方案资料", tool="retrieve") as step:
        outcome = pipeline.retrieve_only(query, k=top_k or settings.retrieval_top_k)
        step.summary = f"命中 {len(outcome.items)} 段方案资料"
        step.detail = "、".join(item.document_name for item in outcome.items[:6]) or "无命中"
        step.status = "success" if outcome.items else "warning"
        step.outputs = {
            "query": query[:200],
            "mode": outcome.stats.mode,
            "after_rerank": outcome.stats.after_rerank,
        }

    response = compose_solution(
        requirement=requirement_text,
        analysis=analysis,
        chunks=outcome.items,
        settings=settings,
        trace=trace,
    )
    response.latency_ms = trace.elapsed_ms
    response.trace = trace.steps
    return response


# ---------------------------------------------------------------------------
# Export
# ---------------------------------------------------------------------------


def markdown_to_docx(markdown: str, title: str, target_path) -> str:
    """Render the proposal to a .docx file.  Returns the written path."""
    try:
        import docx
        from docx.shared import Pt
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("缺少 python-docx 依赖，无法导出 Word。") from exc

    from pathlib import Path

    document = docx.Document()
    document.add_heading(title, level=0)

    def add_inline(paragraph, text: str) -> None:
        """Render **bold** spans inside a paragraph."""
        for position, piece in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
            run = paragraph.add_run(piece)
            if position % 2 == 1:
                run.bold = True

    in_code = False
    for raw_line in (markdown or "").splitlines():
        line = raw_line.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            paragraph = document.add_paragraph(stripped)
            paragraph.style = document.styles["Normal"]
            paragraph.paragraph_format.left_indent = Pt(18)
            continue
        if not stripped:
            continue

        if stripped.startswith("#"):
            level = min(len(stripped) - len(stripped.lstrip("#")), 4)
            document.add_heading(stripped[level:].strip(), level=level)
            continue
        if stripped.startswith(("- ", "* ", "+ ")):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, stripped[2:])
            continue
        if stripped.startswith("|"):
            # Tables are flattened into a monospaced-ish line: the proposal's
            # value is the prose, and a half-rendered table is worse than a
            # readable row.
            if set(stripped) <= set("|-: "):
                continue
            paragraph = document.add_paragraph(stripped.strip("|").replace("|", "  ·  "))
            continue
        if stripped.startswith(">"):
            paragraph = document.add_paragraph()
            run = paragraph.add_run(stripped.lstrip("> ").strip())
            run.italic = True
            continue

        add_inline(document.add_paragraph(), stripped)

    path = Path(target_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return str(path)
